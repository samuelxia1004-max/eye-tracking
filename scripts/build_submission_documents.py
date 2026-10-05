"""Build APA-style DOCX submission files from the reviewed Markdown sources."""

from __future__ import annotations

import html
import os
import re
from pathlib import Path


PROJECT = Path(__file__).resolve().parents[1]
MANUSCRIPT = PROJECT / "manuscript"
SUBMISSION = PROJECT / "submission"
EQUATIONS = SUBMISSION / "_equations"
os.environ.setdefault("MPLCONFIGDIR", str(PROJECT / ".mplconfig"))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "Crossed randomization distinguishes global ocular-score differences from event-boundary-specific effects in naturalistic eye tracking"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def enable_field_updates(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_document(document: Document, *, double_spaced: bool = True) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    fmt = normal.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE if double_spaced else WD_LINE_SPACING.SINGLE
    fmt.widow_control = True

    for name, size, bold, italic, alignment in (
        ("Title", 14, True, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 1", 12, True, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 12, True, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12, True, True, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.alignment = alignment
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True

    if "Equation" not in document.styles:
        equation = document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        equation = document.styles["Equation"]
    equation.font.name = "Cambria Math"
    equation.font.size = Pt(11)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(6)
    equation.paragraph_format.space_after = Pt(6)
    equation.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.style = document.styles["Normal"]
    header.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_field(header, "PAGE", "1")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = footer.add_run("Dual-null inference for event-boundary eye tracking")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(96, 96, 96)

    document.core_properties.title = TITLE
    document.core_properties.subject = "Behavior Research Methods methodology manuscript"
    document.core_properties.comments = "Local submission package; author metadata remain to be completed."
    enable_field_updates(document)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend((color, underline))
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def math_to_unicode(value: str) -> str:
    value = re.sub(r"\\operatorname\{([^}]+)\}", r"\1", value)
    value = re.sub(r"\\bar\{?([A-Za-z])\}?", lambda m: m.group(1) + "\u0304", value)
    value = re.sub(r"\\sqrt\{([^}]+)\}", r"√(\1)", value)
    value = re.sub(r"\\frac\{([^}]+)\}\{([^}]+)\}", r"(\1)/(\2)", value)
    replacements = {
        r"\theta": "θ",
        r"\Delta": "Δ",
        r"\alpha": "α",
        r"\beta": "β",
        r"\mu": "μ",
        r"\sigma": "σ",
        r"\rho": "ρ",
        r"\leq": "≤",
        r"\geq": "≥",
        r"\pm": "±",
        r"\times": "×",
        r"\ldots": "…",
        r"\max": "max",
        r"\left": "",
        r"\right": "",
        r"\!": "",
        r"\qquad": "  ",
        r"\mathbf": "",
        r"\boldsymbol": "",
        r"\mathsf": "",
        r"\mathrm": "",
        r"\text": "",
        r"\mbox": "",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = value.replace("{", "").replace("}", "")
    return value


INLINE_TOKEN = re.compile(
    r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`|\[[^]]+\]\(https?://[^)]+\)|https?://\S+|\\\(.*?\\\)|\$[^$]+\$)"
)


def add_inline(paragraph, text: str) -> None:
    text = html.unescape(text)
    text = re.sub(r"<sub>(.*?)</sub>", r"_\1", text)
    pos = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif token.startswith("["):
            label, url = re.match(r"\[([^]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        elif token.startswith("http"):
            trailing = token.rstrip(".,;)")
            add_hyperlink(paragraph, trailing, trailing)
            if trailing != token:
                paragraph.add_run(token[len(trailing) :])
        else:
            content = token[2:-2] if token.startswith(r"\(") else token[1:-1]
            run = paragraph.add_run(math_to_unicode(content))
            run.font.name = "Cambria Math"
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def render_equation(tex: str, index: int) -> Path | None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    cleaned = tex.strip().rstrip(".,")
    cleaned = cleaned.replace(r"\operatorname{atanh}", r"\mathrm{atanh}")
    cleaned = cleaned.replace(r"\operatorname{cor}", r"\mathrm{cor}")
    cleaned = re.sub(r"\\operatorname\{([^}]+)\}", r"\\mathrm{\1}", cleaned)
    cleaned = cleaned.replace(r"\mbox", r"\mathrm")
    cleaned = cleaned.replace(r"\text", r"\mathrm")
    cleaned = re.sub(r"\\mathbf\s+([A-Za-z])", r"\\mathbf{\1}", cleaned)
    cleaned = re.sub(r"\\mathsf\s+([A-Za-z])", r"\\mathsf{\1}", cleaned)
    cleaned = " ".join(line.strip() for line in cleaned.splitlines())
    output = EQUATIONS / f"equation_{index:02d}.png"
    try:
        fig = plt.figure(figsize=(6.25, 0.55), dpi=300)
        fig.patch.set_alpha(0)
        fig.text(0.5, 0.5, f"${cleaned}$", ha="center", va="center", fontsize=11, color="black")
        fig.savefig(output, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.04)
        plt.close(fig)
        return output
    except Exception:
        plt.close("all")
        return None


def add_equation(document: Document, tex: str, index: int) -> None:
    path = render_equation(tex, index)
    paragraph = document.add_paragraph(style="Equation")
    if path is None:
        run = paragraph.add_run(math_to_unicode(" ".join(tex.splitlines())))
        run.font.name = "Cambria Math"
    else:
        run = paragraph.add_run()
        shape = run.add_picture(str(path), width=Inches(5.9))
        shape._inline.docPr.set("descr", f"Displayed equation {index}")


def parse_table(lines: list[str]):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows


def add_table(document: Document, lines: list[str]) -> None:
    rows = parse_table(lines)
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.margin_top = Inches(0.03)
            value = row[c_idx] if c_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.5)
                if r_idx == 0:
                    run.bold = True
        if r_idx == 0:
            set_repeat_table_header(table.rows[0])
            for cell in table.rows[0].cells:
                set_cell_shading(cell, "D9E2F3")
    document.add_paragraph()


def add_toc(document: Document) -> None:
    title = document.add_paragraph("Contents", style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    entries = [
        ("Introduction", 3),
        ("Dual-null framework", 6),
        ("    One estimand, two null targets", 6),
        ("    Conjunctive decision and evidence states", 8),
        ("    External benchmark and scope", 9),
        ("Method", 10),
        ("    Study 1: BBBD distracted-session benchmark", 11),
        ("    Study 2: Frozen transport to Sherlock", 15),
        ("    Simulation", 18),
        ("    Reproducibility", 19),
        ("Results", 20),
        ("Discussion", 25),
        ("Declarations", 32),
        ("References", 35),
        ("Tables", 39),
        ("Figures", 41),
    ]
    for label, page in entries:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        paragraph.add_run(f"{label}\t{page}")
    document.add_page_break()


def add_markdown(document: Document, source: Path, *, include_toc: bool = False) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    equation_index = 0
    toc_added = False
    in_references = False
    in_code = False
    code_lines: list[str] = []
    figure_index = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                paragraph.paragraph_format.left_indent = Inches(0.25)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(8.5)
                in_code = False
                code_lines = []
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped == r"\[":
            equation: list[str] = []
            i += 1
            while i < len(lines) and lines[i].strip() != r"\]":
                equation.append(lines[i])
                i += 1
            equation_index += 1
            add_equation(document, "\n".join(equation), equation_index)
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(document, table_lines)
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            alt, relative = image_match.groups()
            image_path = (source.parent / relative).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            shape = paragraph.add_run().add_picture(str(image_path), width=Inches(6.4))
            shape._inline.docPr.set("descr", alt)
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            in_references = level == 1 and title == "References"
            if level == 1 and title in {"References", "Tables", "Figures"}:
                document.add_page_break()
            if level == 2 and title.startswith("Figure "):
                figure_index += 1
                if figure_index > 1:
                    document.add_page_break()
            if level == 1 and document.paragraphs == []:
                paragraph = document.add_paragraph(style="Title")
                add_inline(paragraph, title)
            else:
                paragraph = document.add_paragraph(style=f"Heading {level}")
                add_inline(paragraph, title)
            i += 1
            continue
        if stripped.startswith(('- ', '* ')):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            add_inline(paragraph, stripped[2:])
            i += 1
            continue
        paragraph = document.add_paragraph()
        if source.name == "references_apa.md" or in_references:
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        add_inline(paragraph, stripped)
        if stripped.startswith("**Keywords:**") and include_toc and not toc_added:
            add_toc(document)
            toc_added = True
        i += 1


def save_main() -> Path:
    document = Document()
    configure_document(document)
    add_markdown(document, MANUSCRIPT / "manuscript_anonymized.md", include_toc=True)
    output = SUBMISSION / "BRM_manuscript_anonymized.docx"
    document.save(output)
    return output


def save_supplement() -> Path:
    document = Document()
    configure_document(document, double_spaced=False)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    document.styles["Normal"].paragraph_format.space_after = Pt(4)
    add_markdown(document, MANUSCRIPT / "supplement.md")
    output = SUBMISSION / "BRM_supplement.docx"
    document.save(output)
    return output


def save_title_page() -> Path:
    document = Document()
    configure_document(document)
    add_markdown(document, MANUSCRIPT / "title_page.md")
    output = SUBMISSION / "BRM_title_page_template.docx"
    document.save(output)
    return output


def save_cover_letter() -> Path:
    document = Document()
    configure_document(document, double_spaced=False)
    add_markdown(document, MANUSCRIPT / "cover_letter_brm.md")
    output = SUBMISSION / "BRM_cover_letter.docx"
    document.save(output)
    return output


def main() -> None:
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    EQUATIONS.mkdir(parents=True, exist_ok=True)
    outputs = [save_main(), save_supplement(), save_title_page(), save_cover_letter()]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
